"""Build the submission slide deck (16:9 PowerPoint) with speaker notes.

Example:
    python scripts/build_deck.py --output docs/report/AutoAnalog-RL_deck.pptx --team "Team Name"
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "docs" / "report" / "figures"
SHOTS = ROOT / "reports" / "screenshots"

NAVY = RGBColor(0x16, 0x3A, 0x6B)
BLUE = RGBColor(0x2B, 0x57, 0x97)
GREEN = RGBColor(0x1A, 0x7F, 0x37)
RED = RGBColor(0xB4, 0x23, 0x18)
GREY = RGBColor(0x55, 0x5B, 0x66)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

W, H = Inches(13.333), Inches(7.5)


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    # ---- primitives -------------------------------------------------------
    def slide(self, title: str, notes: str = "", subtitle: str | None = None):
        s = self.prs.slides.add_slide(self.blank)
        self.count += 1
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.0))
        bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(11.5), Inches(0.8))
        p = tb.text_frame.paragraphs[0]; p.text = title
        p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = WHITE
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        num = s.shapes.add_textbox(Inches(12.3), Inches(0.3), Inches(0.8), Inches(0.5))
        q = num.text_frame.paragraphs[0]; q.text = str(self.count); q.font.size = Pt(12); q.font.color.rgb = WHITE; q.alignment = PP_ALIGN.RIGHT
        if subtitle:
            st = s.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.5))
            r = st.text_frame.paragraphs[0]; r.text = subtitle; r.font.size = Pt(16); r.font.color.rgb = GREY; r.font.italic = True
        foot = s.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(8), Inches(0.35))
        f = foot.text_frame.paragraphs[0]; f.text = "AutoAnalog-RL  ·  Nebula: AI/ML for Analog Circuit Design  ·  Round 0"; f.font.size = Pt(10); f.font.color.rgb = GREY
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def bullets(self, s, items, left, top, width, height, size=18, color=BLACK):
        tb = s.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame; tf.word_wrap = True
        first = True
        for item in items:
            level = 0
            if isinstance(item, tuple):
                item, level = item
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = ("•  " if level == 0 else "–  ") + item
            p.level = level
            p.font.size = Pt(size - 3 * level); p.font.color.rgb = color
            p.space_after = Pt(6)
        return tb

    def text(self, s, txt, left, top, width, height, size=16, color=BLACK, bold=False, mono=False, align=None, fill=None):
        tb = s.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame; tf.word_wrap = True
        if fill is not None:
            tb.fill.solid(); tb.fill.fore_color.rgb = fill
        for i, line in enumerate(txt.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
            if mono:
                p.font.name = "Consolas"
            if align:
                p.alignment = align
        return tb

    def image(self, s, path, left, top, width=None, height=None):
        return s.shapes.add_picture(str(path), left, top, width=width, height=height)

    def table(self, s, rows, left, top, width, col_widths=None, size=13, header=True, height=None):
        n_rows, n_cols = len(rows), len(rows[0])
        shape = s.shapes.add_table(n_rows, n_cols, left, top, width, height or Inches(0.4) * n_rows)
        t = shape.table
        if col_widths:
            for i, w in enumerate(col_widths):
                t.columns[i].width = w
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c); cell.text = str(val)
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(size)
                    p.font.bold = header and r == 0
                    p.font.color.rgb = WHITE if (header and r == 0) else BLACK
                cell.fill.solid(); cell.fill.fore_color.rgb = BLUE if (header and r == 0) else (LIGHT if r % 2 == 0 else WHITE)
                cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.03)
        return shape

    def stat(self, s, big, small, left, top, width=Inches(2.9), color=NAVY):
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(1.5))
        box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = BLUE
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = big; p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = color; p.alignment = PP_ALIGN.CENTER
        q = tf.add_paragraph(); q.text = small; q.font.size = Pt(12); q.font.color.rgb = GREY; q.alignment = PP_ALIGN.CENTER

    def save(self, path: Path) -> None:
        self.prs.save(str(path))


def build(team: str) -> Presentation:
    d = Deck()

    # 1 Title ---------------------------------------------------------------
    s = d.prs.slides.add_slide(d.blank); d.count += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    d.text(s, "AutoAnalog-RL", Inches(0.8), Inches(1.9), Inches(11.5), Inches(1.2), size=54, color=WHITE, bold=True)
    d.text(s, "Automated Sizing of High-Speed Interface Circuits via Reinforcement Learning", Inches(0.8), Inches(3.0), Inches(11.5), Inches(0.8), size=24, color=WHITE)
    d.text(s, "PCIe Gen 2 (5 Gbps) CTLE + one-tap DFE  ·  IHP sg13g2 130 nm open PDK  ·  ngspice in the loop", Inches(0.8), Inches(3.9), Inches(11.5), Inches(0.6), size=16, color=RGBColor(0xC8, 0xD4, 0xE8))
    d.text(s, f"{team}\nNebula — AI/ML for Analog Circuit Design  ·  Round 0  ·  15 September 2026", Inches(0.8), Inches(5.6), Inches(11.5), Inches(1.0), size=16, color=WHITE)
    s.notes_slide.notes_text_frame.text = (
        "Hello. This is AutoAnalog-RL: a reinforcement-learning agent that sizes a PCIe Gen 2 receiver equalizer by driving "
        "ngspice directly, with zero human intervention, on the IHP sg13g2 130 nm open PDK. In the next few minutes I will show "
        "the problem, the system, a live demo, the evidence, and what we found along the way.")

    # 2 Problem -------------------------------------------------------------
    s = d.slide("The problem", notes=(
        "The brief asks for a fully automated framework that sizes an equalizer for a PCIe PHY from target specifications, "
        "reaching near-optimal designs in far less time than sweeping every MOS, R, C and L value. The specs are coupled: the same "
        "transistor width sets gain, bandwidth, power, linearity and noise, and the one that matters most, the eye after the "
        "channel, only shows up in a transient simulation. Manual sweeps do not scale and generic optimisers burn their budget "
        "in infeasible regions."))
    d.bullets(s, [
        "Size an analog equalizer (CTLE + 1-tap DFE) for a PCIe Gen 2 PHY from target specs — zero human intervention",
        "Coupled specs: one width sets gain, bandwidth, power, linearity, noise; the eye is only visible in a transient",
        "Must beat sweeping the MOS / R / C / L space; must hold across 45 PVT corners",
        "Simulator-bound: every candidate costs several ngspice runs",
    ], Inches(0.6), Inches(1.4), Inches(7.2), Inches(4.5), size=18)
    d.table(s, [
        ["Spec (brief)", "Target"],
        ["Nyquist", "2.5 GHz (5 Gbps), peaking tunable 1.25–2.5 GHz"],
        ["HF peaking", "3–12 dB"],
        ["Eye after channel", "> 100 mV, > 0.4 UI"],
        ["HD3 (100 MHz)", "< −30 dB"],
        ["Noise 10 MHz–5 GHz", "< 1.5 mV rms"],
        ["Power / Area", "< 15 mW / < 0.05 mm²"],
        ["PVT", "TT SS FF SF FS · VDD ±5% · 0–125 °C"],
    ], Inches(8.1), Inches(1.5), Inches(4.8), col_widths=[Inches(1.9), Inches(2.9)], size=12)

    # 3 System --------------------------------------------------------------
    s = d.slide("The system: RL agent → fail-fast SPICE gates → sized netlist", notes=(
        "Here is the loop. The SAC agent proposes a normalised design vector. The environment maps it to device values and the "
        "evaluator runs the ngspice gates in fail-fast order: DC operating point first, then AC for peaking and power, then a "
        "PRBS7 transient through a lossy channel for the eye, then HD3 once everything else passes. Anything that fails early "
        "costs nothing more. The reward is transparent, and an LLM frontend can re-weight it from plain-language feedback. "
        "Eight ngspice processes run in parallel on a laptop: about 0.6 seconds per RL step on the real PSP103 models."))
    boxes = [
        ("SAC agent\nstable-baselines3", NAVY), ("Environment\nΔ actions, reward,\ncorner curriculum", NAVY),
        ("SpiceEvaluator\n1 .op (fail fast)\n2 .ac peaking, power\n3 PRBS7 .tran + channel → eye, DFE\n4 HD3 once others pass", BLUE),
        ("ngspice 47 ×8\nIHP sg13g2 PSP103\nvia OpenVAF / OSDI", NAVY), ("Outputs\nsized netlist\nJSON / CSV / PNG", GREEN),
    ]
    x = Inches(0.4)
    widths = [Inches(1.9), Inches(2.2), Inches(3.4), Inches(2.2), Inches(1.9)]
    gap = Inches(0.3)
    for index, ((label, colr), w) in enumerate(zip(boxes, widths)):
        if index:
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x - gap + Inches(0.03), Inches(2.45), gap - Inches(0.06), Inches(0.3))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = BLUE; arrow.line.fill.background()
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), w, Inches(2.2))
        b.fill.solid(); b.fill.fore_color.rgb = colr; b.line.fill.background()
        tf = b.text_frame; tf.word_wrap = True
        for i, line in enumerate(label.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line; p.font.size = Pt(15 if i == 0 else 12); p.font.bold = i == 0; p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
        x += w + gap
    llm = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.7), Inches(4.25), Inches(2.3), Inches(0.8))
    llm.fill.solid(); llm.fill.fore_color.rgb = RGBColor(0xB7, 0x79, 0x1F); llm.line.fill.background()
    p = llm.text_frame.paragraphs[0]; p.text = "LLM frontend\nfeedback → reward weights"; p.font.size = Pt(12); p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
    up = s.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(3.7), Inches(3.78), Inches(0.3), Inches(0.42))
    up.fill.solid(); up.fill.fore_color.rgb = RGBColor(0xB7, 0x79, 0x1F); up.line.fill.background()
    d.bullets(s, [
        "Bounded 5-parameter action: W_in, R_load, I_bias, R_s, C_s (6 with the DFE tap)",
        "Reward = −Σ wᵢ·violationᵢ − power charge + success bonus + margin bonus; DC failure −10",
        "Real IHP PSP103 models (OpenVAF → OSDI), 8 simulators wide, 0.6 s per environment step",
    ], Inches(0.6), Inches(5.3), Inches(12.2), Inches(1.7), size=14)

    # 4 Circuit ---------------------------------------------------------------
    s = d.slide("The circuit and the specs the reward sees", notes=(
        "The CTLE is a differential NMOS pair with resistive loads and RC source degeneration. The agent sizes five values inside "
        "these bounds. The transient gate sends a PRBS7 at 5 gigabits through a two-section RC channel with about 10 dB of loss at "
        "Nyquist, and the eye is measured after it. We use tighter eye targets than the brief, 250 millivolts and 0.7 UI, so the "
        "policy delivers margin rather than a marginal pass."))
    d.image(s, FIG / "ac_response.png", Inches(0.5), Inches(1.4), width=Inches(6.0))
    d.table(s, [
        ["Knob", "Range"], ["W_in", "0.5–50 µm"], ["R_load", "100–1000 Ω"], ["I_bias", "0.1–2 mA"], ["R_s", "10–500 Ω"], ["C_s", "1 fF–1 pF"], ["DFE tap", "−0.5…+0.5 (optional)"],
    ], Inches(6.9), Inches(1.4), Inches(2.8), col_widths=[Inches(1.1), Inches(1.7)], size=12)
    d.table(s, [
        ["Spec (ours)", "Target", "Gate"],
        ["Peaking @ Nyquist", "3–12 dB", ".ac"], ["Power", "≤ 2 mW", ".op"], ["Eye height", "≥ 0.25 V", "PRBS7 .tran"],
        ["Eye width", "≥ 0.7 UI", "PRBS7 .tran"], ["HD3", "≤ −30 dB", "tone .tran"], ["Noise", "≤ 1.5 mV rms", ".noise (validation)"], ["PVT", "45/45", "5×3×3 corners"],
    ], Inches(9.9), Inches(1.4), Inches(3.1), col_widths=[Inches(1.25), Inches(0.9), Inches(0.95)], size=11)
    d.text(s, "Source-degenerated differential CTLE; devices are IHP sg13_lv_nmos, L = 130 nm. AC response of a policy design: 4.2 dB of peaking at 2.5 GHz.",
           Inches(0.5), Inches(5.0), Inches(6.2), Inches(1.0), size=12, color=GREY)

    # 5 Live demo -------------------------------------------------------------
    s = d.slide("Live: the policy sizes a CTLE from a random start", subtitle="python scripts/demo_rollout.py reports/sac-ihp-pvt   (random design, random PVT corner)", notes=(
        "This is a real rollout. The policy starts from a random design at a random PVT corner, here slow-slow silicon at 62 "
        "degrees. Each line is one ngspice evaluation. Watch the spec checklist: peaking and power pass immediately, the eye is "
        "closed, and within three simulations, seven seconds of wall time, every spec is green including HD3, and the sized "
        "netlist is written. I will run this live in the video with a fresh random start."))
    demo = (FIG / "demo_output.txt").read_text(encoding="utf-8").strip().splitlines()
    demo = [line.replace("  [ok] ", " OK:").replace("  [--] ", " --:").replace("[ok] ", "OK:").replace("[--] ", "--:") for line in demo[:12]]
    box = d.text(s, "\n".join(demo), Inches(0.4), Inches(1.65), Inches(12.5), Inches(3.4), size=9.5, mono=True, fill=RGBColor(0x0C, 0x1B, 0x2E), color=RGBColor(0xE6, 0xEE, 0xF8))
    box.text_frame.word_wrap = False
    d.stat(s, "3 sims", "to meet every spec", Inches(0.6), Inches(5.3))
    d.stat(s, "7 s", "wall time, laptop", Inches(3.7), Inches(5.3))
    d.stat(s, "SS / 1.0 V / 62 °C", "random corner, random start", Inches(6.8), Inches(5.3), width=Inches(3.4))

    # 6 Training ---------------------------------------------------------------
    s = d.slide("Training: seven runs on the real models, all converge within ~2k steps", notes=(
        "These are the learning curves for all seven runs on the IHP models: three seeds, HD3 enforced, the PVT curriculum, the "
        "whole equalizer, and the 1.25 gigahertz retune. Reward per step, fraction of steps meeting every spec, and eye height "
        "against the 0.25 volt target. Every run reaches 92 to 97 percent feasible steps within about two thousand steps; the "
        "curriculum and retune stages start converged because they resume a trained policy. Each run is one to two hours on a laptop."))
    d.image(s, FIG / "learning_curves.png", Inches(0.4), Inches(1.4), width=Inches(12.5))
    d.text(s, "SAC, 8 thread-parallel envs, hold-on-success episodes with a margin bonus · 12k steps ≈ 2 h on a 12-core laptop", Inches(0.5), Inches(6.3), Inches(12), Inches(0.5), size=13, color=GREY)

    # 7 Results table -----------------------------------------------------------
    s = d.slide("Results: every policy design passes every measured spec", notes=(
        "Deterministic rollouts from 20 random starting designs per run. Every run: all rollouts reach a fully feasible design, "
        "in a median of two to three simulations, worst case five. The best designs sit at 0.32 to 0.37 volts of eye after the "
        "channel, around one milliwatt, HD3 near minus 60 dB, and 45 out of 45 PVT corners. Three seeds agree to within a few "
        "percent: 2.53 plus or minus 0.06 steps to feasibility."))
    d.table(s, [
        ["Run", "Variant", "Train feasible", "Rollouts", "Steps (med / worst)", "Best design: eye · power · HD3 · PVT"],
        ["seed 1 (30k)", "—", "95.5%", "20/20", "2 / 4", "0.327 V · 1.22 mW · −57.8 dB · 45/45"],
        ["seed 2 (12k)", "—", "95.5%", "20/20", "3 / 4", "0.316 V · 1.22 mW · −57.7 dB · 45/45"],
        ["seed 3 (12k)", "—", "95.5%", "20/20", "2.5 / 4", "0.348 V · 1.33 mW · −60.2 dB · 45/45"],
        ["hd3 (12k)", "HD3 enforced in reward", "95.6%", "20/20", "2.5 / 4", "0.335 V · 1.24 mW · −58.1 dB · 45/45"],
        ["pvt (+8k)", "curriculum: random corner per episode", "95.4%", "24/24 @ random corners", "2 / 4", "0.370 V · 1.12 mW · −61.9 dB · 45/45"],
        ["eq (8k)", "CTLE + DFE tap in the action", "92.3%", "20/20", "3 / 5", "0.325→0.355 V post-DFE · 1.36 mW · −71.6 dB · 45/45"],
        ["gen1 (+3k)", "retuned to 1.25 GHz Nyquist", "96.5%", "20/20", "2 / 5", "5.23 dB @ 1.25 GHz · 0.350 V · 1.28 mW · 45/45"],
    ], Inches(0.4), Inches(1.4), Inches(12.5), col_widths=[Inches(1.4), Inches(2.9), Inches(1.3), Inches(1.9), Inches(1.5), Inches(3.5)], size=12)
    d.text(s, "Three seeds: 60/60 rollouts feasible · 2.53 ± 0.06 steps to feasibility · eye 0.330 ± 0.014 V · power 1.26 ± 0.05 mW · HD3 −58.6 ± 1.2 dB",
           Inches(0.5), Inches(5.6), Inches(12.3), Inches(0.6), size=14, color=NAVY, bold=True)

    # 8 Baselines -------------------------------------------------------------
    s = d.slide("Versus random search and CMA-ES: the win is amortisation", notes=(
        "Two baselines on the same models and reward. Random search satisfies every spec 7.9 percent of the time, first hit at "
        "design 27. CMA-ES is a strong optimiser here: given 400 simulations per design it reaches the same quality the policy "
        "reaches in training, and from a random start it needs a median of nine evaluations to get feasible. The policy needs "
        "two to three. So the honest claim is amortisation: train once, then size any new instance, a new start, a new corner, a "
        "re-weighted reward, in a couple of simulations with no per-instance search."))
    d.image(s, FIG / "budget_curve.png", Inches(0.4), Inches(1.4), width=Inches(6.3))
    d.table(s, [
        ["Method (random starts)", "First feasible (med / worst)", "Best @ 50 sims", "Per new instance"],
        ["SAC policy (144 rollouts)", "2–3 / 5", "20.5–20.8", "2–3 sims, no search"],
        ["CMA-ES, random start", "9 / 24", "20.50", "~10 to feasible, 400 to refine"],
        ["CMA-ES, box centre", "2–5", "20.8–21.1 (21.3 @ 400)", "400 per design"],
        ["Random search", "27", "19.65", "~13 per feasible design"],
    ], Inches(6.9), Inches(1.5), Inches(6.1), col_widths=[Inches(2.1), Inches(1.5), Inches(1.3), Inches(1.2)], size=11)
    d.bullets(s, [
        "Policy: 3–4× fewer simulations than CMA-ES from the same starts, ~10× fewer than random",
        "CMA-ES matches the policy's training-best quality at 400 sims per design — stated honestly",
        "Bounded AC-only search (pre-ML) never reaches the eye spec: 0.160 V",
    ], Inches(6.9), Inches(4.3), Inches(6.2), Inches(2.5), size=13)

    # 9 Eye before/after -----------------------------------------------------------
    s = d.slide("Eye after the −10 dB channel: AC-only search vs. the policy", notes=(
        "Left: the design chosen by the pre-ML bounded search, which only looks at the AC response. Its eye after the channel is "
        "160 millivolts and 0.58 UI, below our target. Right: the policy design, rewarded on the post-channel eye, 327 millivolts "
        "and 0.88 UI. Same circuit, same channel; the difference is that the agent sees the eye."))
    d.image(s, FIG / "eye_diagram_preml.png", Inches(0.4), Inches(1.5), width=Inches(6.2))
    d.image(s, FIG / "eye_diagram.png", Inches(6.8), Inches(1.5), width=Inches(6.2))
    d.text(s, "Bounded search, AC gate only: 0.160 V / 0.58 UI  —  fails", Inches(0.5), Inches(5.7), Inches(6), Inches(0.5), size=15, color=RED, bold=True)
    d.text(s, "SAC policy (seed 1): 0.327 V / 0.88 UI  —  passes with margin", Inches(6.9), Inches(5.7), Inches(6), Inches(0.5), size=15, color=GREEN, bold=True)

    # 10 PVT curriculum -----------------------------------------------------------
    s = d.slide("Curriculum across 45 PVT corners", notes=(
        "Stage one trains at the nominal corner. Stage two resumes the same policy and simulates every episode at a random one "
        "of the 45 corners, five process, three supply, three temperature, without telling the policy which. It transfers "
        "immediately, 95 percent feasible from the first window, and in evaluation 24 out of 24 rollouts at random corners reach "
        "feasibility in a median of two steps. The best design passes all 45 corners with peaking well inside the window."))
    d.image(s, FIG / "pvt_peaking.png", Inches(0.4), Inches(1.4), width=Inches(8.4))
    d.stat(s, "24 / 24", "rollouts feasible at random corners", Inches(9.2), Inches(1.6), width=Inches(3.7))
    d.stat(s, "45 / 45", "corners pass, best design", Inches(9.2), Inches(3.3), width=Inches(3.7))
    d.stat(s, "2 steps", "median to feasibility", Inches(9.2), Inches(5.0), width=Inches(3.7))
    d.text(s, "Per process corner during training: SS 97% · FS 96% · FF 95% · TT 95% · SF 94%", Inches(0.5), Inches(5.6), Inches(8.4), Inches(0.5), size=13, color=GREY)

    # 11 Tunable + equalizer -------------------------------------------------------
    s = d.slide("Specs in → schematic out: tunable Nyquist and the DFE tap", notes=(
        "The brief asks for peaking tunable from 1.25 to 2.5 gigahertz. The Nyquist frequency is a specification input: it sets "
        "the measurement frequency, the data rate, and the channel pole. The 2.5 gigahertz design only peaks 2.87 dB at 1.25, so it "
        "fails there; resuming the policy with the spec set to 1.25 gigahertz gives 20 out of 20 rollouts and 5.23 dB, and the "
        "agent moves the degeneration zero down in frequency exactly as a designer would. Second, the whole equalizer: with the "
        "DFE tap in the action, the policy learns to spend the tap on margin, holding a half-volt post-DFE eye during training."))
    d.text(s, "Tunable Nyquist frequency", Inches(0.5), Inches(1.4), Inches(6), Inches(0.5), size=20, color=NAVY, bold=True)
    d.bullets(s, [
        "--spec nyquist_frequency_hz=1.25e9 → PCIe Gen 1 (2.5 Gbps): re-derives measurement frequency, UI, transient step, channel pole",
        "2.5 GHz design at 1.25 GHz: 2.87 dB — fails",
        "Resume policy, 3k steps: 96.5% feasible, 20/20 rollouts, median 2 steps",
        "Best design: 5.23 dB @ 1.25 GHz, 0.350 V / 0.94 UI, 1.28 mW, HD3 −70.9 dB, 45/45",
        "R_s 182 → 365 Ω, C_s 522 → 503 fF: the zero moves down in frequency",
    ], Inches(0.5), Inches(2.0), Inches(6.2), Inches(4.5), size=14)
    d.text(s, "Whole equalizer: CTLE + DFE tap", Inches(6.9), Inches(1.4), Inches(6), Inches(0.5), size=20, color=NAVY, bold=True)
    d.image(s, FIG / "eye_diagram_equalizer.png", Inches(7.4), Inches(1.95), height=Inches(3.6))
    d.text(s, "6-value action; post-DFE eye (measured against the transmitted bits) drives the reward. Training holds 0.51 V post-DFE vs ~0.35 V CTLE-only; rollouts 20/20, tap +0.062 opens 0.325 → 0.355 V with zero bit errors.",
           Inches(6.9), Inches(5.7), Inches(6.0), Inches(1.2), size=12, color=GREY)

    # 12 LLM frontend --------------------------------------------------------------
    s = d.slide("LLM frontend: plain-language feedback → reward weights", notes=(
        "The bonus item. An engineer types what they care about; the script turns it into the reward's per-spec weights, the "
        "power charge and the margin bonus. Claude produces them as structured JSON when credentials exist; a context-aware keyword "
        "parser is the offline fallback so the loop never blocks. The JSON is passed to training with --reward-settings and "
        "applies from the start of that run or resumed stage."))
    d.text(s, '> python scripts/reward_from_feedback.py "We are power constrained: current budget matters much more\n   than eye margin, and we don\'t care about linearity"\n\n'
              '{ "weights": { "power": 3.0, "eye_vertical_v": 0.33, "hd3": 0.33, "peaking_boost": 1.0, ... },\n'
              '  "efficiency_weight": 3.0, "margin_weight": 5.0,\n'
              '  "rationale": "raised power x3; lowered eye_vertical_v x0.33; lowered hd3 x0.33" }\n\n'
              '> python scripts/train_sac.py ... --reward-settings reports/reward_low_power.json',
           Inches(0.5), Inches(1.5), Inches(12.3), Inches(3.0), size=13, mono=True, fill=RGBColor(0x0C, 0x1B, 0x2E), color=RGBColor(0xE6, 0xEE, 0xF8))
    d.bullets(s, [
        "Claude (Anthropic SDK, schema-validated JSON) when credentials are present; keyword parser offline — the loop never blocks on an API",
        "Reward: −Σ wᵢ·vᵢ − w_eff·(P/P_max) + [feasible]·(bonus + w_margin·margin) — every wᵢ is what the frontend adjusts",
        "Applies at the start of a run or a resumed curriculum stage; re-weighting mid-run is the next step",
    ], Inches(0.5), Inches(4.8), Inches(12.3), Inches(2.0), size=14)

    # 13 Bugs found ---------------------------------------------------------------
    s = d.slide("What the agent found: two measurement flaws the manual flow lived with", notes=(
        "Two things worth being candid about. First, HD3 had been reported as failing at minus 25 dB in every earlier version of "
        "this project. That was spectral leakage: the FFT analysed two and a half cycles with a rectangular window. With an "
        "integer-cycle Hann window the same designs measure minus 58 dB, 28 dB of margin. The agent exposed it by being unable "
        "to move the number. Second, the DFE eye was labelled by the DFE's own decisions, so a large tap could fake a wide eye; "
        "the equalizer agent found that in two thousand steps and pushed the tap to its bound while the receiver got 54 of 127 "
        "bits wrong. The eye is now measured against the transmitted bits. Reward hacking as a bug detector."))
    d.table(s, [
        ["Flaw", "Symptom", "Cause", "Fix", "After"],
        ["HD3 measurement", "−25 dB 'fail' on every design; RL could not move it past −26", "2.5-cycle rectangular FFT window → fundamental leakage in the 300 MHz bin", "integer-cycle Hann window (unit-tested on a synthetic −40 dB tone)", "−58 to −72 dB, 28 dB margin"],
        ["DFE eye metric", "agent drove the tap to +0.5 and reported a 0.70 V eye", "eye labelled by the DFE's own decisions — the tap separates the classes by itself", "eye against the transmitted PRBS bits; bit errors counted; zero eye on any error", "54/127 errors exposed; honest DFE still +39% eye"],
    ], Inches(0.4), Inches(1.5), Inches(12.5), col_widths=[Inches(1.6), Inches(2.8), Inches(3.1), Inches(3.0), Inches(2.0)], size=12)
    d.text(s, "Both fixed, unit-tested, and every number in the report uses the corrected measurements. Noise was cross-checked against ngspice's own band integral (0.2420 vs 0.2417 mV rms).",
           Inches(0.5), Inches(3.6), Inches(12.3), Inches(0.9), size=14, color=NAVY)
    d.stat(s, "−25 → −58 dB", "HD3, same design, correct window", Inches(0.6), Inches(4.8), width=Inches(3.9))
    d.stat(s, "54 / 127", "bit errors behind the 'best' DFE eye", Inches(4.7), Inches(4.8), width=Inches(3.9))
    d.stat(s, "+39 %", "honest one-tap DFE eye gain", Inches(8.8), Inches(4.8), width=Inches(3.9), color=GREEN)
    d.text(s, "Reward hacking as a bug detector: an optimiser that can exploit a metric will find its flaws first.",
           Inches(0.5), Inches(6.45), Inches(12.3), Inches(0.5), size=15, color=GREY, bold=True, align=PP_ALIGN.CENTER)

    # 14 Compliance -----------------------------------------------------------------
    s = d.slide("Against the brief: every specification, every deliverable", notes=(
        "Mapping the brief line by line. Nyquist, peaking with tunability, the CTLE with variable degeneration, the one-tap DFE, "
        "HD3, noise, power, area, eye, PVT, zero human intervention and faster than sweeping, the Python framework with specs in "
        "and schematic out, and the LLM bonus: all delivered. The one change from our own synopsis is the PDK: IHP instead of "
        "sky130, which the brief allows; we switched because the sky130 flow was too slow for an RL loop."))
    d.table(s, [
        ["Brief", "Ours", "Result", ""],
        ["Nyquist 2.5 GHz, peaking 3–12 dB tunable 1.25–2.5 GHz", "spec input", "4.5 dB @ 2.5 GHz; 5.23 dB @ 1.25 GHz", "✓"],
        ["1-stage CTLE, variable Rs/Cs; 1-tap DFE; NRZ", "in the action", "yes; tap sized by the agent", "✓"],
        ["HD3 < −30 dB", "enforced", "−61.9 dB", "✓"],
        ["Noise < 1.5 mV rms", "reported", "0.242 mV rms", "✓"],
        ["Power < 15 mW · Area < 0.05 mm²", "≤ 2 mW · reported", "1.12 mW · 1.1e-5 mm² (estimate)", "✓"],
        ["Eye > 100 mV, > 0.4 UI", "≥ 0.25 V, ≥ 0.7 UI", "0.370 V, 0.90 UI", "✓"],
        ["PVT TT/SS/FF/SF/FS, ±5%, 0–125 °C", "trained across + verified", "45/45", "✓"],
        ["Zero human intervention, faster than sweeping", "2–3 sims from a random start", "random 13, CMA-ES 9", "✓"],
        ["Python RL framework, specs in, SPICE in loop, schematic out", "--spec, ngspice gates, sized_ctle.sp", "yes", "✓"],
        ["Bonus: LLM human interaction", "reward_from_feedback.py", "yes", "✓"],
        ["Open PDK: IHP 130 nm or sky130", "IHP sg13g2 (synopsis said sky130)", "PSP103 via OSDI, 0.6 s/step", "✓"],
    ], Inches(0.4), Inches(1.4), Inches(12.5), col_widths=[Inches(4.6), Inches(3.2), Inches(3.9), Inches(0.8)], size=12)

    # 15 Deliverables + next -----------------------------------------------------------
    s = d.slide("Deliverables, limitations, next steps", notes=(
        "The deliverables: the repository with 59 tests, the 12-page report, every run's artifacts including the sized netlist, "
        "and the live demo. Limitations we state openly: the DFE in the loop is behavioural, area is a first-order estimate, and "
        "the curriculum policy is corner-blind. Next steps in order of value: the analog slicer inside the transient gate, "
        "corner-aware observations, a closed LLM loop, and a sky130 backend if a future round needs it. Thank you."))
    d.text(s, "Deliverables", Inches(0.5), Inches(1.4), Inches(4), Inches(0.5), size=20, color=NAVY, bold=True)
    d.bullets(s, [
        "github.com/Foolish-Genius/al_nebula — 59 tests, reproducible commands",
        "Report (12 pp, PDF + LaTeX) with every number traceable to an artifact",
        "sized_ctle.sp + validation.json for every design; checkpoints; 45-corner PVT",
        "scripts/demo_rollout.py — the live demo",
    ], Inches(0.5), Inches(2.0), Inches(4.2), Inches(4.5), size=13)
    d.text(s, "Limitations", Inches(4.9), Inches(1.4), Inches(4), Inches(0.5), size=20, color=NAVY, bold=True)
    d.bullets(s, [
        "DFE in the loop is behavioural; analog slicer not simulated end to end",
        "Area: first-order device estimate, not post-layout",
        "Curriculum policy is corner-blind (robust, not corner-aware)",
        "PDK: IHP, not the synopsis' sky130 (allowed by the brief)",
    ], Inches(4.9), Inches(2.0), Inches(4.0), Inches(4.5), size=13)
    d.text(s, "Next steps", Inches(9.1), Inches(1.4), Inches(4), Inches(0.5), size=20, color=NAVY, bold=True)
    d.bullets(s, [
        "Analog DFE slicer inside the transient gate",
        "Corner-aware observation",
        "Closed LLM loop: train → read report → re-weight → resume",
        "BO baseline + wall-clock vs. manual sweep",
        "sky130 backend if required",
    ], Inches(9.1), Inches(2.0), Inches(4.0), Inches(4.5), size=13)
    d.text(s, "Thank you — questions?", Inches(0.5), Inches(6.2), Inches(12), Inches(0.6), size=22, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    return d.prs


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default="docs/report/AutoAnalog-RL_deck.pptx")
    parser.add_argument("--team", default="Team: TBD")
    args = parser.parse_args()
    prs = build(args.team)
    out = ROOT / args.output
    prs.save(str(out))
    print(f"wrote {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
